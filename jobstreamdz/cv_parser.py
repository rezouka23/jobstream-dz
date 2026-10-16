from __future__ import annotations

from io import BytesIO
from pathlib import Path


SUPPORTED_CV_EXTENSIONS = {".pdf", ".docx", ".txt"}


class CVParseError(ValueError):
    pass


def supported_cv_extensions_label() -> str:
    return ", ".join(sorted(SUPPORTED_CV_EXTENSIONS))


def extract_cv_text(*, filename: str, data: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension not in SUPPORTED_CV_EXTENSIONS:
        raise CVParseError(f"Unsupported CV format. Use {supported_cv_extensions_label()}.")

    if extension == ".pdf":
        text = _extract_pdf_text(data)
    elif extension == ".docx":
        text = _extract_docx_text(data)
    else:
        text = _extract_txt_text(data)

    normalized = _normalize_cv_text(text)
    if not normalized:
        raise CVParseError("The CV text could not be read. Please upload a text-based PDF, DOCX, or TXT file.")
    return normalized


def _extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # pragma: no cover - dependency failure path
        raise CVParseError("PDF support is not available on the server.") from exc

    try:
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise CVParseError("The PDF could not be read. Please upload a text-based PDF.") from exc


def _extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - dependency failure path
        raise CVParseError("DOCX support is not available on the server.") from exc

    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except Exception as exc:
        raise CVParseError("The DOCX file could not be read.") from exc


def _extract_txt_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CVParseError("The text file could not be decoded.")


def _normalize_cv_text(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    return "\n".join(line for line in lines if line).strip()
